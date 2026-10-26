import io

from docx import Document
from google.genai import types

from .llm import generate_text

_MAX_FILE_BYTES = 10 * 1024 * 1024
_CHUNK_LIMIT = 6000
_MAX_CHUNKS = 20
_SUMMARIZE_LIMIT = 30000
_TASKS = ("translate", "summarize")


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith((".txt", ".md")):
        return data.decode("utf-8", errors="replace")
    if name.endswith(".docx"):
        document = Document(io.BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(parts)
    raise ValueError("Format file tidak didukung.")


def _instruction(task: str, target_language: str) -> str:
    if task == "translate":
        return (
            f"Translate the document into {target_language}. "
            "Keep structure and paragraphs. Output only the translation."
        )
    return (
        f"Summarize the document in {target_language} as bullet points (max 15). "
        "Output only the summary."
    )


def _split_chunks(text: str, limit: int = _CHUNK_LIMIT) -> list[str]:
    chunks = []
    current = []
    for paragraph in text.split("\n\n"):
        trial = "\n\n".join(current + [paragraph]) if current else paragraph
        if len(trial) <= limit:
            current.append(paragraph)
        elif current:
            chunks.append("\n\n".join(current))
            current = [paragraph]
        else:
            chunks.append(paragraph)
            current = []
    if current:
        chunks.append("\n\n".join(current))
    return chunks


def _generate(parts: list, instruction: str) -> str:
    return generate_text([{"role": "user", "parts": parts}], instruction, timeout_ms=60000)


def process_document(
    filename: str, data: bytes, mime_type: str, task: str, target_language: str
) -> str:
    if len(data) > _MAX_FILE_BYTES:
        raise RuntimeError("File terlalu besar (maks 10 MB).")
    if task not in _TASKS:
        raise ValueError("Task tidak valid.")
    instruction = _instruction(task, target_language)
    name = filename.lower()
    if name.endswith((".pdf", ".png", ".jpg", ".jpeg", ".webp")):
        parts = [types.Part.from_bytes(data=data, mime_type=mime_type), {"text": instruction}]
        return _generate(parts, instruction)
    text = extract_text(filename, data)
    if task == "summarize":
        if len(text) > _SUMMARIZE_LIMIT:
            raise RuntimeError("Dokumen terlalu panjang")
        return _generate([{"text": text}], instruction)
    if len(text) > _CHUNK_LIMIT:
        chunks = _split_chunks(text)
        if len(chunks) > _MAX_CHUNKS:
            raise RuntimeError("Dokumen terlalu panjang")
        return "\n\n".join(_generate([{"text": chunk}], instruction) for chunk in chunks)
    return _generate([{"text": text}], instruction)


def to_docx_bytes(title: str, text: str) -> bytes:
    document = Document()
    document.core_properties.title = title
    for line in text.splitlines():
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()